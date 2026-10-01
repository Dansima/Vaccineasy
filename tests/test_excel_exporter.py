from datetime import datetime
from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from app.business_logic import get_all_vaccination_statuses
from app.excel_exporter import convert_df_to_catagrafie
from app.reporting import filter_monthly_report


def make_row(dob, ref, name='Copil Test', cnp='5250101123456'):
    pending = get_all_vaccination_statuses(dob, ref)
    return {'Nume si Prenume': name, 'CNP': cnp, 'Vârsta_datetime': dob,
            'Status': pending[0][0] if pending else '🟢 La Zi', '_pending': pending,
            '_all_codes': [c for _, _, c in pending], 'Vaccin Necesar': ''}


def test_mixed_dose_statuses_and_totals():
    ref = datetime(2025, 5, 1)
    frame = pd.DataFrame([make_row(datetime(2025, 1, 1), ref, '=HYPERLINK("x")')])
    data = convert_df_to_catagrafie(frame, ref)
    book = load_workbook(BytesIO(data))
    sheet = book.active
    assert sheet['B10'].data_type == 's' and sheet['C10'].data_type == 's'
    assert sheet['E10'].value == sheet['K10'].value == 'X'  # Hexa 2 overdue
    assert sheet['F10'].value == sheet['L10'].value == 'X'  # Hexa 4 due
    assert sheet['G10'].value is None
    assert sheet.freeze_panes == 'D10'
    assert 'luna 5 / anul 2025' in sheet['A5'].value
    assert sheet['E23'].value == '=COUNTIF(E10:E22,"X")'
    assert sheet.page_setup.orientation == 'landscape'
    values = load_workbook(BytesIO(data), data_only=True).active
    assert values['E23'].value == 1 and values['F23'].value == 1
    assert values['C24'].value == 1


def test_month_filter_excludes_next_month_dose_even_for_overdue_patient():
    ref = datetime(2025, 4, 15)
    frame = pd.DataFrame([make_row(datetime(2025, 1, 1), ref)])
    filtered = filter_monthly_report(frame, ref)
    assert filtered.iloc[0]['_all_codes'] == ['Hexa_2']
    sheet = load_workbook(BytesIO(convert_df_to_catagrafie(frame, ref))).active
    assert sheet['E10'].value == 'X'
    assert sheet['F10'].value is None and sheet['G10'].value is None


def test_pagination_sorting_and_global_totals():
    ref = datetime(2025, 5, 31)
    frame = pd.DataFrame([make_row(datetime(2025, 1, day), ref, f'Copil {day}') for day in range(1, 15)])
    book = load_workbook(BytesIO(convert_df_to_catagrafie(frame, ref)), data_only=True)
    assert book.sheetnames == ['Pagina 1', 'Pagina 2']
    first, second = book.worksheets
    assert first['B10'].value == 'Copil 14'
    assert second['B10'].value == 'Copil 1'
    assert second['A10'].value == 14
    assert first['C23'].value == 13 and second['C23'].value == 1
    assert first['C24'].value == second['C24'].value == 14
    assert first['E24'].value == second['E24'].value == 14


def test_empty_export():
    frame = pd.DataFrame(columns=['Status', '_pending', 'Vârsta_datetime'])
    sheet = load_workbook(BytesIO(convert_df_to_catagrafie(frame)), data_only=True).active
    assert sheet['C23'].value == sheet['C24'].value == 0


def test_legacy_single_code_export_has_correct_cached_totals():
    frame = pd.DataFrame([{'Nume si Prenume': 'Test', 'CNP': '5250101123456',
                           'Status': '🔴 RESTANT', '_cod_cat': 'Hexa_2'}])
    sheet = load_workbook(BytesIO(convert_df_to_catagrafie(frame)), data_only=True).active
    assert sheet['E10'].value == sheet['K10'].value == 'X'
    assert sheet['E23'].value == sheet['E24'].value == 1
